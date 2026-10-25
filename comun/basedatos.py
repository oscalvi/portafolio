"""
Manejo genérico de datos, en programación orientada a objetos.

ESTRUCTURA
    GestorBD                 Clase padre de lo relacional: abrir, cerrar, ejecutar, 'with'.
    ├── Conexion             Cómo hablar con cada motor (una subclase por motor):
    │     ├── SQLite             conectar, marcador de parámetros, tipos de campo,
    │     ├── Access             autonumérico, id del registro nuevo, límite de filas,
    │     ├── SQLServer          crear la base vacía, listar tablas.
    │     ├── MySQL
    │     └── Oracle
    └── Administracion       Estructura de la base (DDL): crear base y tablas,
                             borrar tablas, agregar campos, vistas, procedimientos.

    Almacen                  INTERFAZ ÚNICA para los registros, sea cual sea el tipo:
    ├── Datos                    bases relacionales (usa una Conexion)
    └── DatosArchivo             archivos, sin servidor:
          ├── DatosJSON              una lista JSON por tabla
          ├── DatosCSV               archivo secuencial: un .csv por tabla
          └── DatosExcel             un libro de Excel, una hoja por tabla

    La aplicación siempre usa los mismos métodos (leer, buscar, contar, escribir,
    actualizar, borrar, transaccion) sin saber qué hay detrás. Para agregar otro
    tipo (por ejemplo MongoDB) basta con escribir otra subclase de Almacen.

REGLAS QUE SE CUMPLEN SIEMPRE
    1. El TIPO de base de datos va siempre en los parámetros: la fábrica lo exige.
    2. actualizar y borrar exigen el ID del registro. Nunca se modifican ni borran
       registros "en general".
    3. escribir devuelve el ID del registro nuevo (para enlazar factura -> detalle).
    4. Los nombres de tablas y campos se validan antes de usarlos en SQL; los valores
       viajan siempre como parámetros (protección contra inyección SQL).
    5. Si el diseño (tablas) es conocido, cada registro se valida: campos que no existen,
       obligatorios vacíos, tipo de dato y largo máximo.
    6. 'with almacen.transaccion():' confirma todo junto o deshace todo si algo falla.

LIBRERÍA NECESARIA POR TIPO
    "sqlite"     -> incluida en Python
    "sqlserver"  -> pip install pyodbc   (y ODBC Driver 17 for SQL Server)
    "access"     -> pip install pyodbc pywin32   (y driver de Access)
    "oracle"     -> pip install oracledb
    "mysql"      -> pip install mysql-connector-python
    "json", "csv" -> incluidas en Python
    "excel"      -> pip install openpyxl

PARÁMETROS DE CONEXIÓN POR TIPO
    sqlite     -> archivo
    access     -> archivo (.accdb o .mdb)
    sqlserver  -> servidor, base, y usuario + clave  (o confianza=True para Windows)
    mysql      -> servidor, base, usuario, clave
    oracle     -> usuario, clave, dsn  (ej. "servidor:1521/servicio")
    json, csv  -> carpeta
    excel      -> archivo (.xlsx)

USO
    from basedatos import Almacen

    datos = Almacen.desde_config("bd.json")      # el "tipo" del JSON decide la clase
    with datos:
        id_ana = datos.escribir("clientes", {"nombre": "Ana"})       # devuelve el id
        datos.actualizar("clientes", id_ana, {"nombre": "Ana María"})
        print(datos.buscar("clientes", id_ana))
        print(datos.leer("clientes", {"nombre LIKE": "Ana%"}, orden="-id", limite=10))
        datos.borrar("clientes", id_ana)

        with datos.transaccion():                # todo o nada
            id_f = datos.escribir("facturas", {"cliente_id": 1, "valor": 100})
            datos.escribir("detalles", {"factura_id": id_f, "producto": "Café"})

    Condiciones de leer/contar: {"campo": valor} (igual), o con operador en la clave:
        {"fecha >=": "2026-01-01", "valor <": 500, "nombre LIKE": "%ana%", "estado !=": "X"}
    orden: "nombre" (ascendente) o "-fecha" (descendente), o una lista de ellos.

    Estructura (solo bases relacionales):
        from basedatos import Conexion, Administracion
        con = Conexion.desde_config("bd.json")
        with Administracion(con) as admin:
            admin.crear_esquema(con.config)      # crea lo que falte

FORMATO DEL ARCHIVO DE CONFIGURACIÓN (bd.json)
    {
      "tipo": "sqlite",
      "conexion": {"archivo": "clientes.db"},
      "tablas": {
        "clientes": {
          "id":       {"tipo": "entero", "pk": true, "auto": true},
          "nombre":   {"tipo": "texto", "largo": 60, "obligatorio": true}
        },
        "facturas": {
          "id":         {"tipo": "entero", "pk": true, "auto": true},
          "cliente_id": {"tipo": "entero", "obligatorio": true},
          "valor":      {"tipo": "decimal"}
        }
      },
      "relaciones": [
        {"tabla": "facturas", "campo": "cliente_id", "referencia": "clientes.id"}
      ],
      "vistas":         {"v_facturas": "SELECT ... FROM ..."},
      "procedimientos": {"sp_nombre": "<cuerpo SQL propio del motor>"}
    }
    Tipos genéricos de campo: entero, texto, decimal, fecha.
    "auto": true -> el motor numera solo (IDENTITY, AUTO_INCREMENT, COUNTER...).
    "relaciones", "vistas" y "procedimientos" son opcionales.
"""
import csv
import json
import re
import sqlite3
from abc import ABC, abstractmethod
from contextlib import contextmanager
from datetime import date, datetime
from pathlib import Path

try:                                        # mensajes en el idioma del visitante (si existe)
    from idiomas import t
except ImportError:                         # usado fuera de la aplicación web
    def t(texto, **datos):
        return texto.format(**datos) if datos else texto


class ErrorDatos(ValueError):
    """Error de validación: el mensaje se puede mostrar al usuario tal cual."""


# Nombres de tablas y campos: letra o _ al inicio, luego letras, números o _
NOMBRE_VALIDO = re.compile(r"[A-Za-z_][A-Za-z0-9_]{0,63}")
OPERADORES = {"=", "!=", "<>", "<", "<=", ">", ">=", "LIKE"}


def validar_nombre(nombre):
    """Devuelve el nombre si es seguro para usarlo en SQL; si no, ErrorDatos."""
    if not isinstance(nombre, str) or not NOMBRE_VALIDO.fullmatch(nombre):
        raise ErrorDatos(t("Nombre no válido: «{nombre}». Use letras, números o _ "
                           "(sin espacios ni tildes).", nombre=nombre))
    return nombre


def separar_condicion(clave):
    """'fecha >=' -> ('fecha', '>='); 'nombre' -> ('nombre', '=')."""
    partes = clave.split()
    campo = validar_nombre(partes[0])
    operador = " ".join(partes[1:]).upper() or "="
    if operador not in OPERADORES:
        raise ErrorDatos(t("Operador no permitido: «{op}».", op=operador))
    return campo, ("<>" if operador == "!=" else operador)


def separar_orden(orden):
    """'-fecha' -> [('fecha', True)]; ['nombre', '-id'] -> [('nombre', False), ('id', True)]."""
    if not orden:
        return []
    lista = [orden] if isinstance(orden, str) else list(orden)
    return [(validar_nombre(o.lstrip("-")), o.startswith("-")) for o in lista]


# =====================================================================
#  CLASE PADRE (bases relacionales)
# =====================================================================

class GestorBD:
    """Comportamiento común: todo pasa por un objeto Conexion."""

    def __init__(self, conexion):
        self.con = conexion

    def abrir(self):
        self.con.abrir()
        return self

    def cerrar(self):
        self.con.cerrar()

    def ejecutar(self, sql, valores=()):
        return self.con.ejecutar(sql, valores)

    def confirmar(self):
        self.con.confirmar()

    # abrir y cerrar solos con 'with'; si hubo un error, se deshace lo pendiente
    def __enter__(self):
        return self.abrir()

    def __exit__(self, tipo_error, *resto):
        if tipo_error and self.con.enlace:
            self.con.deshacer()
        self.cerrar()


# =====================================================================
#  CONEXIÓN: lo que cambia según el motor
# =====================================================================

class Conexion(GestorBD):
    """Clase base de los motores. Cada subclase define cómo conectarse."""

    TIPOS = {"entero": "INTEGER", "texto": "VARCHAR({largo})",
             "decimal": "DECIMAL(12,2)", "fecha": "DATE"}
    AUTO = "INTEGER PRIMARY KEY"            # cómo se escribe un id autonumérico
    soporta_procedimientos = True

    def __init__(self, config=None, **parametros):
        super().__init__(self)              # la conexión es su propio gestor
        self.parametros = parametros
        self.config = config or {}
        self.enlace = None                  # conexión real del driver
        self._usos = 0                      # permite 'with' anidados

    # ---------- fábrica ----------
    MOTORES = {}                            # se llena al final de las subclases

    @staticmethod
    def crear(tipo, config=None, **parametros):
        if not tipo:
            raise ErrorDatos(t("Falta el tipo de base de datos."))
        try:
            return Conexion.MOTORES[tipo.lower()](config, **parametros)
        except KeyError:
            raise ErrorDatos(t("Tipo de base de datos no soportado: {tipo}", tipo=tipo)) from None

    @staticmethod
    def desde_config(ruta_json):
        with open(ruta_json, encoding="utf-8") as f:
            config = json.load(f)
        return Conexion.crear(config.get("tipo"), config, **config["conexion"])

    # ---------- abrir y cerrar ----------
    @abstractmethod
    def conectar(self):
        """Devuelve la conexión del driver. Cada motor la implementa."""

    def abrir(self):
        if self._usos == 0:
            self.enlace = self.conectar()
        self._usos += 1
        return self

    def cerrar(self):
        self._usos = max(self._usos - 1, 0)
        if self._usos == 0 and self.enlace:
            self.enlace.close()
            self.enlace = None

    def probar(self):
        """Abre y cierra: lanza un error si algo está mal."""
        with self:
            pass

    def ejecutar(self, sql, valores=()):
        cursor = self.enlace.cursor()
        cursor.execute(sql, list(valores))
        return cursor

    def confirmar(self):
        self.enlace.commit()

    def deshacer(self):
        self.enlace.rollback()

    # ---------- dialecto: lo que cada motor escribe distinto ----------
    def marcador(self, n):
        """Marcador de parámetro en las instrucciones SQL."""
        return "?"

    def tipo_campo(self, definicion):
        if definicion.get("auto"):
            return self.AUTO
        return self.TIPOS[definicion["tipo"]].format(largo=definicion.get("largo", 50))

    def sql_limite(self, select, n):
        """Limita las filas de un SELECT. SQLite y MySQL: LIMIT n."""
        return f"{select} LIMIT {int(n)}"

    def insertar(self, tabla, datos, campo_id):
        """INSERT y devuelve el id del registro nuevo. SQLite y MySQL: lastrowid."""
        cursor = self.ejecutar(self._sql_insert(tabla, datos), datos.values())
        return cursor.lastrowid

    def _sql_insert(self, tabla, datos, extra=""):
        columnas = ", ".join(datos)
        marcas = ", ".join(self.marcador(i + 1) for i in range(len(datos)))
        return f"INSERT INTO {tabla} ({columnas}){extra} VALUES ({marcas})"

    def sql_procedimiento(self, nombre, cuerpo):
        return f"CREATE PROCEDURE {nombre} AS {cuerpo}"

    def crear_base(self):
        raise NotImplementedError(f"En {type(self).__name__} la base la crea el "
                                  "administrador; pídale los datos de conexión.")

    def listar_tablas(self):
        cursor = self.ejecutar("SELECT TABLE_NAME FROM INFORMATION_SCHEMA.TABLES "
                               "WHERE TABLE_TYPE = 'BASE TABLE'")
        return sorted(fila[0] for fila in cursor.fetchall())


class SQLite(Conexion):
    """Parámetros: archivo"""
    TIPOS = {**Conexion.TIPOS, "texto": "TEXT", "decimal": "REAL", "fecha": "TEXT"}
    AUTO = "INTEGER PRIMARY KEY"
    soporta_procedimientos = False

    def conectar(self):
        enlace = sqlite3.connect(self.parametros["archivo"])
        enlace.execute("PRAGMA foreign_keys = ON")
        return enlace

    def crear_base(self):
        archivo = Path(self.parametros["archivo"])
        if archivo.exists():
            raise FileExistsError(f"Ya existe {archivo}.")
        archivo.parent.mkdir(parents=True, exist_ok=True)
        self.probar()                       # SQLite crea el archivo al conectarse

    def listar_tablas(self):
        cursor = self.ejecutar("SELECT name FROM sqlite_master WHERE type = 'table' "
                               "AND name NOT LIKE 'sqlite_%' ORDER BY name")
        return [fila[0] for fila in cursor.fetchall()]


class Access(Conexion):
    """Parámetros: archivo (.accdb o .mdb)"""
    TIPOS = {"entero": "LONG", "texto": "TEXT({largo})",
             "decimal": "CURRENCY", "fecha": "DATETIME"}
    AUTO = "COUNTER PRIMARY KEY"
    soporta_procedimientos = False

    def conectar(self):
        import pyodbc
        return pyodbc.connect(r"DRIVER={Microsoft Access Driver (*.mdb, *.accdb)};DBQ="
                              + self.parametros["archivo"])

    def crear_base(self):
        import win32com.client              # pip install pywin32
        archivo = Path(self.parametros["archivo"])
        if archivo.exists():
            raise FileExistsError(f"Ya existe {archivo}.")
        archivo.parent.mkdir(parents=True, exist_ok=True)
        catalogo = win32com.client.Dispatch("ADOX.Catalog")
        catalogo.Create(f"Provider=Microsoft.ACE.OLEDB.12.0;Data Source={archivo}")

    def listar_tablas(self):
        return sorted(t.table_name for t in self.enlace.cursor().tables(tableType="TABLE"))

    def sql_limite(self, select, n):
        return select.replace("SELECT ", f"SELECT TOP {int(n)} ", 1)

    def insertar(self, tabla, datos, campo_id):
        self.ejecutar(self._sql_insert(tabla, datos), datos.values())
        return self.ejecutar("SELECT @@IDENTITY").fetchone()[0]


class SQLServer(Conexion):
    """Parámetros: servidor, base, usuario, clave (o confianza=True para Windows)"""
    TIPOS = {**Conexion.TIPOS, "texto": "NVARCHAR({largo})"}
    AUTO = "INT IDENTITY(1,1) PRIMARY KEY"

    def _cadena(self, base):
        p = self.parametros
        cadena = (f"DRIVER={{ODBC Driver 17 for SQL Server}};SERVER={p['servidor']};"
                  f"DATABASE={base};")
        if p.get("confianza"):
            return cadena + "Trusted_Connection=yes;"
        return cadena + f"UID={p['usuario']};PWD={p['clave']};"

    def conectar(self):
        import pyodbc
        return pyodbc.connect(self._cadena(self.parametros["base"]))

    def crear_base(self):
        import pyodbc
        with pyodbc.connect(self._cadena("master"), autocommit=True) as enlace:
            enlace.execute(f"CREATE DATABASE [{validar_nombre(self.parametros['base'])}]")

    def sql_limite(self, select, n):
        return select.replace("SELECT ", f"SELECT TOP {int(n)} ", 1)

    def insertar(self, tabla, datos, campo_id):
        # OUTPUT INSERTED devuelve el id generado en la misma instrucción
        cursor = self.ejecutar(self._sql_insert(tabla, datos, f" OUTPUT INSERTED.{campo_id}"),
                               datos.values())
        return cursor.fetchone()[0]


class MySQL(Conexion):
    """Parámetros: servidor, base, usuario, clave"""
    AUTO = "INT AUTO_INCREMENT PRIMARY KEY"

    def conectar(self, con_base=True):
        import mysql.connector
        p = self.parametros
        extra = {"database": p["base"]} if con_base else {}
        return mysql.connector.connect(host=p["servidor"], user=p["usuario"],
                                       password=p["clave"], **extra)

    def crear_base(self):
        enlace = self.conectar(con_base=False)
        enlace.cursor().execute(f"CREATE DATABASE `{validar_nombre(self.parametros['base'])}`")
        enlace.close()

    def listar_tablas(self):
        cursor = self.ejecutar("SELECT TABLE_NAME FROM information_schema.TABLES "
                               "WHERE TABLE_SCHEMA = DATABASE() AND TABLE_TYPE = 'BASE TABLE' "
                               "ORDER BY TABLE_NAME")
        return [fila[0] for fila in cursor.fetchall()]

    def marcador(self, n):
        return "%s"

    def sql_procedimiento(self, nombre, cuerpo):
        return f"CREATE PROCEDURE {nombre}() BEGIN {cuerpo} END"


class Oracle(Conexion):
    """Parámetros: usuario, clave, dsn (ej. 'servidor:1521/servicio')"""
    TIPOS = {"entero": "NUMBER(10)", "texto": "VARCHAR2({largo})",
             "decimal": "NUMBER(12,2)", "fecha": "DATE"}
    AUTO = "NUMBER GENERATED BY DEFAULT AS IDENTITY PRIMARY KEY"

    def conectar(self):
        import oracledb
        p = self.parametros
        return oracledb.connect(user=p["usuario"], password=p["clave"], dsn=p["dsn"])

    def listar_tablas(self):
        cursor = self.ejecutar("SELECT table_name FROM user_tables ORDER BY table_name")
        return [fila[0] for fila in cursor.fetchall()]

    def marcador(self, n):
        return f":{n}"

    def sql_limite(self, select, n):
        return f"{select} FETCH FIRST {int(n)} ROWS ONLY"

    def insertar(self, tabla, datos, campo_id):
        # RETURNING ... INTO: Oracle entrega el id generado en una variable de salida
        cursor = self.enlace.cursor()
        nuevo = cursor.var(int)
        sql = self._sql_insert(tabla, datos) + f" RETURNING {campo_id} INTO :{len(datos) + 1}"
        cursor.execute(sql, list(datos.values()) + [nuevo])
        return nuevo.getvalue()[0]

    def sql_procedimiento(self, nombre, cuerpo):
        return f"CREATE OR REPLACE PROCEDURE {nombre} AS {cuerpo}"


Conexion.MOTORES = {"sqlite": SQLite, "sqlserver": SQLServer, "access": Access,
                    "oracle": Oracle, "mysql": MySQL}


# =====================================================================
#  ADMINISTRACIÓN: estructura de la base (DDL)
# =====================================================================

class Administracion(GestorBD):

    def crear_base(self):
        """Crea la base vacía y verifica que se pueda abrir."""
        self.con.crear_base()
        self.con.probar()

    def listar_tablas(self):
        return self.con.listar_tablas()

    def columnas(self, tabla):
        cursor = self.ejecutar(f"SELECT * FROM {validar_nombre(tabla)} WHERE 1 = 0")
        return [d[0] for d in cursor.description]

    def existe(self, nombre):
        """True si la tabla o vista ya existe."""
        validar_nombre(nombre)
        try:
            self.ejecutar(f"SELECT * FROM {nombre} WHERE 1 = 0")
            return True
        except Exception:
            self.con.deshacer()
            return False

    def _definicion(self, nombre, d):
        col = f"{validar_nombre(nombre)} {self.con.tipo_campo(d)}"
        if d.get("pk") and not d.get("auto"):
            col += " PRIMARY KEY"
        elif d.get("obligatorio"):
            col += " NOT NULL"
        return col

    def crear_tabla(self, tabla, campos, relaciones=()):
        """campos:     {"id": {"tipo": "entero", "pk": True, "auto": True}, ...}
           relaciones: [{"tabla": tabla, "campo": "cliente_id", "referencia": "clientes.id"}]"""
        validar_nombre(tabla)
        columnas = [self._definicion(n, d) for n, d in campos.items()]
        for r in relaciones:
            if r["tabla"] == tabla:
                ref_tabla, ref_campo = r["referencia"].split(".")
                columnas.append(f"FOREIGN KEY ({validar_nombre(r['campo'])}) "
                                f"REFERENCES {validar_nombre(ref_tabla)}({validar_nombre(ref_campo)})")
        self.ejecutar(f"CREATE TABLE {tabla} ({', '.join(columnas)})")
        self.confirmar()

    def borrar_tabla(self, tabla):
        self.ejecutar(f"DROP TABLE {validar_nombre(tabla)}")
        self.confirmar()

    def agregar_campo(self, tabla, nombre, definicion):
        self.ejecutar(f"ALTER TABLE {validar_nombre(tabla)} ADD {self._definicion(nombre, definicion)}")
        self.confirmar()

    def crear_vista(self, nombre, sql):
        """El SQL de la vista lo escribe el administrador en el diseño (bd.json)."""
        self.ejecutar(f"CREATE VIEW {validar_nombre(nombre)} AS {sql}")
        self.confirmar()

    def crear_procedimiento(self, nombre, cuerpo):
        if not self.con.soporta_procedimientos:
            print(f"Aviso: {type(self.con).__name__} no soporta procedimientos almacenados.")
            return
        self.ejecutar(self.con.sql_procedimiento(validar_nombre(nombre), cuerpo))
        self.confirmar()

    def crear_esquema(self, config):
        """Crea lo que falte de un diseño completo (formato de bd.json)."""
        if not config.get("tablas"):
            raise ValueError("Falta la configuración de tablas. "
                             "Ver documentación al inicio de basedatos.py")
        relaciones = config.get("relaciones", [])
        for tabla, campos in config["tablas"].items():
            if not self.existe(tabla):
                self.crear_tabla(tabla, campos, relaciones)
        for nombre, sql in config.get("vistas", {}).items():
            if not self.existe(nombre):
                self.crear_vista(nombre, sql)
        for nombre, cuerpo in config.get("procedimientos", {}).items():
            self.crear_procedimiento(nombre, cuerpo)


# =====================================================================
#  ALMACÉN: la interfaz única para los registros
# =====================================================================

class Almacen(ABC):
    """Lo común a cualquier tipo de almacenamiento: reglas, validación y transacciones.
    Las subclases solo implementan cómo se lee y se guarda en su medio."""

    def __init__(self, tablas=None):
        self.tablas = tablas or {}          # diseño conocido: {"tabla": {"campo": {...}}}
        self._transacciones = 0

    # ---------- fábrica: el tipo decide la clase ----------
    @staticmethod
    def crear(tipo, config=None, **parametros):
        if not tipo:
            raise ErrorDatos(t("Falta el tipo de base de datos."))
        tipo = tipo.lower()
        config = config or {}
        if tipo in ARCHIVOS:
            return ARCHIVOS[tipo](config.get("tablas"), **parametros)
        return Datos(Conexion.crear(tipo, config, **parametros))

    @staticmethod
    def desde_config(ruta_json):
        ruta = Path(ruta_json)
        config = json.loads(ruta.read_text(encoding="utf-8"))
        parametros = dict(config.get("conexion", {}))
        for clave in ("archivo", "carpeta"):  # rutas relativas: junto al bd.json
            if clave in parametros and not Path(parametros[clave]).is_absolute():
                parametros[clave] = str(ruta.parent / parametros[clave])
        return Almacen.crear(config.get("tipo"), config, **parametros)

    # ---------- 'with' ----------
    def abrir(self):
        return self

    def cerrar(self):
        pass

    def __enter__(self):
        return self.abrir()

    def __exit__(self, tipo_error, *resto):
        if tipo_error and self._transacciones == 0:
            self._deshacer()
        self.cerrar()

    # ---------- transacciones: todo o nada ----------
    @contextmanager
    def transaccion(self):
        self.abrir()
        self._transacciones += 1
        try:
            yield self
        except BaseException:
            self._transacciones -= 1
            if self._transacciones == 0:
                self._deshacer()
            self.cerrar()
            raise
        self._transacciones -= 1
        if self._transacciones == 0:
            self._confirmar()
        self.cerrar()

    def _al_cambiar(self):
        """Después de cada cambio: confirma enseguida, salvo dentro de una transacción."""
        if self._transacciones == 0:
            self._confirmar()

    # ---------- reglas comunes ----------
    def campo_id(self, tabla):
        for campo, d in self.tablas.get(tabla, {}).items():
            if d.get("pk"):
                return campo
        return "id"

    @staticmethod
    def _exigir_id(tabla, id_registro):
        if id_registro is None or id_registro == "" or isinstance(id_registro, (dict, list, tuple)):
            raise ErrorDatos(t("Para modificar o borrar en «{tabla}» hay que indicar el ID "
                               "del registro.", tabla=tabla))
        return id_registro

    def _validar(self, tabla, datos, nuevo):
        """Revisa nombres y, si el diseño es conocido, campos, obligatorios, tipos y largos."""
        validar_nombre(tabla)
        if not datos:
            raise ErrorDatos(t("No llegaron datos para guardar en «{tabla}».", tabla=tabla))
        for campo in datos:
            validar_nombre(campo)
        diseno = self.tablas.get(tabla)
        if not diseno:
            return dict(datos)
        limpios = {}
        for campo, valor in datos.items():
            d = diseno.get(campo)
            if d is None:
                raise ErrorDatos(t("La tabla «{tabla}» no tiene el campo «{campo}».",
                                   tabla=tabla, campo=campo))
            if d.get("pk") and not nuevo:
                raise ErrorDatos(t("El ID de un registro no se puede cambiar."))
            limpios[campo] = self._convertir(campo, valor, d)
        if nuevo:
            faltan = [c for c, d in diseno.items() if d.get("obligatorio") and not d.get("auto")
                      and limpios.get(c) in (None, "")]
        else:
            faltan = [c for c in limpios if diseno[c].get("obligatorio") and limpios[c] in (None, "")]
        if faltan:
            raise ErrorDatos(t("Faltan datos obligatorios: {campos}", campos=", ".join(faltan)))
        return limpios

    @staticmethod
    def _convertir(campo, valor, d):
        """Comprueba y convierte el valor según el tipo genérico del campo."""
        if valor is None or valor == "":
            return None if d.get("tipo") != "texto" else valor
        tipo = d.get("tipo", "texto")
        try:
            if tipo == "entero":
                if isinstance(valor, float) and not valor.is_integer():
                    raise ValueError
                return int(valor)
            if tipo == "decimal":
                return float(valor)
            if tipo == "fecha":
                if isinstance(valor, (date, datetime)):
                    return valor.isoformat()[:10]
                return date.fromisoformat(str(valor)[:10]).isoformat()
        except (TypeError, ValueError):
            nombres = {"entero": "un número entero", "decimal": "un número",
                       "fecha": "una fecha AAAA-MM-DD"}
            raise ErrorDatos(t("«{campo}» debe ser {tipo}.", campo=campo,
                               tipo=t(nombres[tipo]))) from None
        texto = str(valor)
        if d.get("largo") and len(texto) > d["largo"]:
            raise ErrorDatos(t("«{campo}» admite máximo {largo} caracteres.",
                               campo=campo, largo=d["largo"]))
        return texto

    # ---------- interfaz pública (igual para todos los tipos) ----------
    def buscar(self, tabla, id_registro):
        """Un registro por su ID, o None si no existe."""
        filas = self.leer(tabla, {self.campo_id(tabla): self._exigir_id(tabla, id_registro)}, limite=1)
        return filas[0] if filas else None

    def escribir(self, tabla, datos):
        """Agrega un registro y devuelve su ID."""
        limpios = self._validar(tabla, datos, nuevo=True)
        nuevo_id = self._insertar(tabla, limpios, self.campo_id(tabla))
        self._al_cambiar()
        return nuevo_id

    def actualizar(self, tabla, id_registro, datos):
        """Modifica UN registro, indicado por su ID."""
        self._exigir_id(tabla, id_registro)
        limpios = self._validar(tabla, datos, nuevo=False)
        if not self._modificar(tabla, self.campo_id(tabla), id_registro, limpios):
            raise ErrorDatos(t("No existe en «{tabla}» un registro con ID {id}.",
                               tabla=tabla, id=id_registro))
        self._al_cambiar()

    def borrar(self, tabla, id_registro):
        """Borra UN registro, indicado por su ID."""
        validar_nombre(tabla)
        self._exigir_id(tabla, id_registro)
        if not self._eliminar(tabla, self.campo_id(tabla), id_registro):
            raise ErrorDatos(t("No existe en «{tabla}» un registro con ID {id}.",
                               tabla=tabla, id=id_registro))
        self._al_cambiar()

    @abstractmethod
    def leer(self, tabla, condiciones=None, campos="*", orden=None, limite=None):
        """Lista de registros (diccionarios)."""

    @abstractmethod
    def contar(self, tabla, condiciones=None):
        """Cuántos registros cumplen las condiciones."""

    # ---------- lo que cada medio implementa ----------
    @abstractmethod
    def _insertar(self, tabla, datos, campo_id):
        """Guarda y devuelve el ID nuevo."""

    @abstractmethod
    def _modificar(self, tabla, campo_id, id_registro, datos):
        """True si encontró y modificó el registro."""

    @abstractmethod
    def _eliminar(self, tabla, campo_id, id_registro):
        """True si encontró y borró el registro."""

    @abstractmethod
    def _confirmar(self):
        """Hace permanentes los cambios."""

    @abstractmethod
    def _deshacer(self):
        """Descarta los cambios no confirmados."""


# =====================================================================
#  DATOS: registros en bases relacionales (DML)
# =====================================================================

class Datos(Almacen, GestorBD):
    """Almacén sobre una Conexion. El diseño se toma de la configuración de la conexión."""

    def __init__(self, conexion):
        GestorBD.__init__(self, conexion)
        Almacen.__init__(self, conexion.config.get("tablas"))

    def abrir(self):
        GestorBD.abrir(self)
        return self

    def cerrar(self):
        GestorBD.cerrar(self)

    def __enter__(self):
        return self.abrir()

    def __exit__(self, tipo_error, *resto):
        Almacen.__exit__(self, tipo_error, *resto)

    def _donde(self, condiciones, inicio=0):
        if not condiciones:
            return "", []
        partes = []
        for i, clave in enumerate(condiciones):
            campo, operador = separar_condicion(clave)
            partes.append(f"{campo} {operador} {self.con.marcador(inicio + i + 1)}")
        return " WHERE " + " AND ".join(partes), list(condiciones.values())

    @staticmethod
    def _campos(campos):
        if campos == "*":
            return "*"
        lista = campos.split(",") if isinstance(campos, str) else campos
        return ", ".join(validar_nombre(c.strip()) for c in lista)

    def leer(self, tabla, condiciones=None, campos="*", orden=None, limite=None):
        where, valores = self._donde(condiciones)
        sql = f"SELECT {self._campos(campos)} FROM {validar_nombre(tabla)}{where}"
        ordenes = separar_orden(orden)
        if ordenes:
            sql += " ORDER BY " + ", ".join(f"{c} DESC" if desc else c for c, desc in ordenes)
        if limite:
            sql = self.con.sql_limite(sql, limite)
        cursor = self.ejecutar(sql, valores)
        columnas = [d[0] for d in cursor.description]
        return [dict(zip(columnas, fila)) for fila in cursor.fetchall()]

    def contar(self, tabla, condiciones=None):
        where, valores = self._donde(condiciones)
        sql = f"SELECT COUNT(*) FROM {validar_nombre(tabla)}{where}"
        return self.ejecutar(sql, valores).fetchone()[0]

    def _insertar(self, tabla, datos, campo_id):
        nuevo_id = self.con.insertar(tabla, datos, validar_nombre(campo_id))
        return datos.get(campo_id, nuevo_id)    # si el ID lo puso el usuario, ese es

    def _modificar(self, tabla, campo_id, id_registro, datos):
        sets = ", ".join(f"{c} = {self.con.marcador(i + 1)}" for i, c in enumerate(datos))
        sql = (f"UPDATE {tabla} SET {sets} "
               f"WHERE {validar_nombre(campo_id)} = {self.con.marcador(len(datos) + 1)}")
        return self.ejecutar(sql, list(datos.values()) + [id_registro]).rowcount > 0

    def _eliminar(self, tabla, campo_id, id_registro):
        sql = f"DELETE FROM {tabla} WHERE {validar_nombre(campo_id)} = {self.con.marcador(1)}"
        return self.ejecutar(sql, [id_registro]).rowcount > 0

    def _confirmar(self):
        self.con.confirmar()

    def _deshacer(self):
        if self.con.enlace:
            self.con.deshacer()


# =====================================================================
#  DATOS EN ARCHIVOS: JSON, CSV (secuencial) y Excel
# =====================================================================

class DatosArchivo(Almacen):
    """Lo común a los archivos: las tablas se cargan en memoria, se filtran aquí mismo
    y se guardan al confirmar. Dentro de una transacción no se escribe nada hasta el final."""

    def __init__(self, tablas=None, **parametros):
        super().__init__(tablas)
        self.parametros = parametros
        self._memoria = {}                  # tabla -> lista de registros
        self._cambiadas = set()

    # --- lo que cada formato implementa ---
    @abstractmethod
    def _cargar(self, tabla):
        """Lista de registros guardados (vacía si la tabla aún no existe)."""

    @abstractmethod
    def _guardar(self, tabla, filas):
        """Escribe todos los registros de la tabla."""

    # --- memoria ---
    def _filas(self, tabla):
        validar_nombre(tabla)
        if tabla not in self._memoria:
            self._memoria[tabla] = [self._tipar(tabla, f) for f in self._cargar(tabla)]
        return self._memoria[tabla]

    def _tipar(self, tabla, fila):
        """Los formatos de texto (CSV, a veces Excel) traen todo como texto: se convierte."""
        diseno = self.tablas.get(tabla, {})
        convertida = {}
        for campo, valor in fila.items():
            d = diseno.get(campo)
            try:
                convertida[campo] = (self._convertir(campo, valor, d)
                                     if d and valor not in (None, "") else valor)
            except ErrorDatos:
                convertida[campo] = valor
        return convertida

    @staticmethod
    def _cumple(fila, condiciones):
        for clave, esperado in (condiciones or {}).items():
            campo, operador = separar_condicion(clave)
            valor = fila.get(campo)
            if operador == "LIKE":
                patron = re.escape(str(esperado)).replace("%", ".*").replace("_", ".")
                if valor is None or not re.fullmatch(patron, str(valor), re.I):
                    return False
                continue
            try:
                ok = {"=": valor == esperado, "<>": valor != esperado,
                      "<": valor is not None and valor < esperado,
                      "<=": valor is not None and valor <= esperado,
                      ">": valor is not None and valor > esperado,
                      ">=": valor is not None and valor >= esperado}[operador]
            except TypeError:
                ok = False
            if not ok:
                return False
        return True

    def leer(self, tabla, condiciones=None, campos="*", orden=None, limite=None):
        filas = [dict(f) for f in self._filas(tabla) if self._cumple(f, condiciones)]
        for campo, desc in reversed(separar_orden(orden)):
            filas.sort(key=lambda f: (f.get(campo) is None, f.get(campo)), reverse=desc)
        if limite:
            filas = filas[:int(limite)]
        if campos != "*":
            lista = campos.split(",") if isinstance(campos, str) else campos
            nombres = [validar_nombre(c.strip()) for c in lista]
            filas = [{c: f.get(c) for c in nombres} for f in filas]
        return filas

    def contar(self, tabla, condiciones=None):
        return sum(1 for f in self._filas(tabla) if self._cumple(f, condiciones))

    def _insertar(self, tabla, datos, campo_id):
        filas = self._filas(tabla)
        diseno = self.tablas.get(tabla, {}).get(campo_id, {})
        if datos.get(campo_id) in (None, ""):
            if diseno and not diseno.get("auto"):
                raise ErrorDatos(t("Falta el ID «{campo}» del registro.", campo=campo_id))
            ids = [f.get(campo_id) for f in filas if isinstance(f.get(campo_id), int)]
            datos = {campo_id: max(ids, default=0) + 1, **datos}
        elif any(f.get(campo_id) == datos[campo_id] for f in filas):
            raise ErrorDatos(t("Ya existe en «{tabla}» un registro con ID {id}.",
                               tabla=tabla, id=datos[campo_id]))
        filas.append(datos)
        self._cambiadas.add(tabla)
        return datos[campo_id]

    def _modificar(self, tabla, campo_id, id_registro, datos):
        for fila in self._filas(tabla):
            if fila.get(campo_id) == id_registro:
                fila.update(datos)
                self._cambiadas.add(tabla)
                return True
        return False

    def _eliminar(self, tabla, campo_id, id_registro):
        filas = self._filas(tabla)
        for i, fila in enumerate(filas):
            if fila.get(campo_id) == id_registro:
                del filas[i]
                self._cambiadas.add(tabla)
                return True
        return False

    def _confirmar(self):
        for tabla in sorted(self._cambiadas):
            self._guardar(tabla, self._memoria[tabla])
        self._cambiadas.clear()

    def _deshacer(self):
        for tabla in self._cambiadas:       # se olvida lo que no se guardó
            self._memoria.pop(tabla, None)
        self._cambiadas.clear()

    def _columnas(self, tabla, filas):
        """Orden de columnas: el del diseño, y luego cualquier otra que traigan los datos."""
        columnas = list(self.tablas.get(tabla, {}))
        for f in filas:
            columnas += [c for c in f if c not in columnas]
        return columnas


class DatosJSON(DatosArchivo):
    """Parámetros: carpeta. Cada tabla es <carpeta>/<tabla>.json (una lista de registros)."""

    def _ruta(self, tabla):
        return Path(self.parametros["carpeta"]) / f"{tabla}.json"

    def _cargar(self, tabla):
        ruta = self._ruta(tabla)
        return json.loads(ruta.read_text(encoding="utf-8")) if ruta.exists() else []

    def _guardar(self, tabla, filas):
        ruta = self._ruta(tabla)
        ruta.parent.mkdir(parents=True, exist_ok=True)
        ruta.write_text(json.dumps(filas, ensure_ascii=False, indent=2), encoding="utf-8")


class DatosCSV(DatosArchivo):
    """Archivo secuencial. Parámetros: carpeta. Cada tabla es <carpeta>/<tabla>.csv."""

    def _ruta(self, tabla):
        return Path(self.parametros["carpeta"]) / f"{tabla}.csv"

    def _cargar(self, tabla):
        ruta = self._ruta(tabla)
        if not ruta.exists():
            return []
        with ruta.open(encoding="utf-8-sig", newline="") as f:
            return list(csv.DictReader(f))

    def _guardar(self, tabla, filas):
        ruta = self._ruta(tabla)
        ruta.parent.mkdir(parents=True, exist_ok=True)
        with ruta.open("w", encoding="utf-8-sig", newline="") as f:   # utf-8-sig: Excel lo abre bien
            escritor = csv.DictWriter(f, fieldnames=self._columnas(tabla, filas))
            escritor.writeheader()
            escritor.writerows(filas)


class DatosExcel(DatosArchivo):
    """Parámetros: archivo (.xlsx). Cada tabla es una hoja; la fila 1 tiene los nombres."""

    def _libro(self):
        import openpyxl                     # pip install openpyxl
        ruta = Path(self.parametros["archivo"])
        return openpyxl.load_workbook(ruta) if ruta.exists() else None

    def _cargar(self, tabla):
        libro = self._libro()
        if not libro or tabla not in libro.sheetnames:
            return []
        filas = list(libro[tabla].iter_rows(values_only=True))
        if not filas:
            return []
        nombres = [str(n) for n in filas[0]]
        return [dict(zip(nombres, f)) for f in filas[1:] if any(v is not None for v in f)]

    def _guardar(self, tabla, filas):
        import openpyxl
        ruta = Path(self.parametros["archivo"])
        libro = self._libro()
        if libro is None:
            libro = openpyxl.Workbook()
            libro.remove(libro.active)
        if tabla in libro.sheetnames:
            libro.remove(libro[tabla])
        hoja = libro.create_sheet(tabla)
        columnas = self._columnas(tabla, filas)
        hoja.append(columnas)
        for f in filas:
            hoja.append([f.get(c) for c in columnas])
        ruta.parent.mkdir(parents=True, exist_ok=True)
        libro.save(ruta)


ARCHIVOS = {"json": DatosJSON, "csv": DatosCSV, "excel": DatosExcel}
